"""
Импорт цен с FTP-сервера: таблица CSV/XLS/XLSX с колонками артикул (код) и
цена (плюс произвольные другие колонки — производитель, название, остатки,
кратность, описание и т.п., они игнорируются, см. parse_price_table).

Сопоставление — НЕ по Part.offer_id (это внутренний код продавца на Ozon),
а по атрибуту Ozon #9048 «ОЕМ-номер / номер детали» (ARTICLE_ATTRIBUTE_ID из
catalog_rules.py) — тому же полю, что уже используется для поиска по
артикулу на сайте. У одного товара в этом атрибуте может быть сразу
несколько номеров через "; " (сам номер + аналоги/кросс-номера — на
практике встречается по 10-15 штук на товар), поэтому код с FTP ищется как
ТОЧНОЕ совпадение с одним из "; "-разделённых значений, а не по вхождению
подстроки (иначе, например, код "133" ложно совпал бы с "1336659"). Если
один и тот же код с FTP совпал с ОЕМ-номером сразу у нескольких разных
товаров (например, один и тот же аналог указан в карточках двух разных
позиций) — цена не проставляется никому, это считается неоднозначным
совпадением (ambiguous), чтобы не приписать цену не тому товару.

Цена с FTP приоритетнее цены Ozon: для товара, для которого нашлось
однозначное совпадение по коду, Part.price всегда берётся из последнего
FTP-файла, а обычная синхронизация с Ozon (sync_ozon.py) больше не
перезаписывает эту цену — см. проверку `part.ftp_price is None` там.
Приоритет снимается сам, как только товар пропадает из файла на FTP при
одном из следующих импортов (см. блок "stale" в run_ftp_price_sync) — тогда
Ozon снова начинает обновлять ему цену как обычно.
"""
import csv
import ftplib
import io
import logging
import os
import re
import time
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Optional

from catalog_rules import ARTICLE_ATTRIBUTE_ID
from database import SessionLocal, init_db
from models import Part, PartAttribute, SyncLog, SyncSettings
from security import validate_outbound_target
import price_inbox

# Потолок размера скачиваемого прайса и общего времени скачивания — защита от
# подмены файла на «бесконечный»/огромный (исчерпание памяти и диска).
MAX_FTP_FILE_BYTES = int(os.getenv("FTP_MAX_FILE_MB", "50")) * 1024 * 1024
MAX_FTP_DOWNLOAD_SECONDS = int(os.getenv("FTP_MAX_DOWNLOAD_SECONDS", "300"))

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("ftp_price_sync")

MAX_UNMATCHED_SAMPLE = 20


def _now():
    return datetime.now(timezone.utc)


def get_or_create_sync_settings(db) -> SyncSettings:
    """Та же строка настроек (id=1), что читает и sync_ozon.py — продублировано
    здесь, чтобы вызов селфчека/FTP-импорта не тянул за собой ozonapi/httpx."""
    settings = db.query(SyncSettings).filter_by(id=1).one_or_none()
    if settings is None:
        settings = SyncSettings(id=1)
        db.add(settings)
        db.commit()
        db.refresh(settings)
    return settings


def _parse_price(raw) -> Optional[Decimal]:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    # Убираем всё, кроме цифр/точки/запятой/минуса — валюта, пробелы между
    # разрядами и т.п. не мешают.
    text = re.sub(r"[^\d,.\-]", "", text)
    if not text:
        return None
    if "," in text and "." not in text:
        text = text.replace(",", ".")  # запятая как десятичный разделитель
    else:
        text = text.replace(",", "")   # запятая — разделитель тысяч
    try:
        value = Decimal(text)
    except InvalidOperation:
        return None
    if value <= 0:
        return None
    return value


def _rows_from_csv(raw_bytes: bytes):
    text = None
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw_bytes.decode("utf-8", errors="replace")

    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") >= sample.count(",") else ","

    reader = csv.reader(io.StringIO(text), dialect)
    return [row for row in reader if any(cell.strip() for cell in row)]


def _normalize_cell(cell):
    """Excel/xlrd отдают ячейку-число как int/float, даже если в файле она
    оформлена как текст с ведущими нулями (типичная проблема с артикулами
    вида "0155305363" — они открываются как число 155305363, ведущий ноль
    теряется безвозвратно на стороне Excel, тут это не восстановить). Здесь
    убираем только артефакт, который вносим мы сами при чтении: превращение
    целого числа в "155305363.0" из-за того, что openpyxl/xlrd отдают такие
    ячейки как float."""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return cell


def _rows_from_xlsx(raw_bytes: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    sheet = wb.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        if any(cell not in (None, "") for cell in row):
            rows.append(["" if c is None else _normalize_cell(c) for c in row])
    return rows


def _rows_from_xls(raw_bytes: bytes):
    import xlrd
    book = xlrd.open_workbook(file_contents=raw_bytes)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = [_normalize_cell(c) for c in sheet.row_values(r)]
        if any(str(c).strip() for c in row):
            rows.append(row)
    return rows


# Заголовки колонок с артикулом и ценой, которые встречаются в реальных
# прайс-листах поставщика (например, колонки «Производитель / Код / Название /
# Остатки / Цена / Кратность / Описание товара» — код и цена не обязательно
# первые две колонки). Сравнение регистронезависимое, по точному совпадению
# заголовка ячейки.
ARTICLE_HEADER_NAMES = {"код", "артикул", "sku", "offer_id", "код товара", "артикул товара"}
PRICE_HEADER_NAMES = {"цена", "price", "цена, руб", "цена руб", "стоимость"}


def _detect_columns_by_header(header_row) -> Optional[tuple]:
    """Возвращает (article_col, price_col) по названиям заголовков, либо None,
    если строка не похожа на заголовок или в ней нет обеих нужных колонок."""
    normalized = [str(cell).strip().lower() for cell in header_row]
    article_col = next((i for i, cell in enumerate(normalized) if cell in ARTICLE_HEADER_NAMES), None)
    price_col = next((i for i, cell in enumerate(normalized) if cell in PRICE_HEADER_NAMES), None)
    if article_col is None or price_col is None:
        return None
    return article_col, price_col


def parse_price_table(filename: str, raw_bytes: bytes):
    """Список (article: str, price: Decimal) для распознанных строк, всё
    остальное (производитель, название, остатки, кратность, описание и т.п.)
    отбрасывается уже на этом этапе.

    Колонки с артикулом и ценой ищутся по заголовку первой строки (см.
    ARTICLE_HEADER_NAMES/PRICE_HEADER_NAMES) — в реальных прайс-листах они не
    обязательно первые две колонки. Если заголовок не распознан (например,
    файл без строки заголовков или с непривычными названиями), используется
    прежнее поведение: первая колонка — артикул, вторая — цена, а строка
    считается заголовком и пропускается, если во второй колонке не похоже на
    цену."""
    lower = (filename or "").lower()
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        rows = _rows_from_xlsx(raw_bytes)
    elif lower.endswith(".xls"):
        rows = _rows_from_xls(raw_bytes)
    else:
        rows = _rows_from_csv(raw_bytes)

    if not rows:
        return []

    start = 0
    columns = _detect_columns_by_header(rows[0])
    if columns is not None:
        article_col, price_col = columns
        start = 1
    else:
        article_col, price_col = 0, 1
        if len(rows[0]) >= 2 and _parse_price(rows[0][1]) is None:
            start = 1  # похоже на строку заголовков — пропускаем

    result = []
    for row in rows[start:]:
        if len(row) <= max(article_col, price_col):
            continue
        article = str(row[article_col]).strip()
        price = _parse_price(row[price_col])
        if not article or price is None:
            continue
        result.append((article, price))
    return result


def _connect_ftp(host, port, user, password, use_tls) -> ftplib.FTP:
    # Проверяем адрес назначения (SSRF) и подключаемся к уже проверенному IP.
    # ftplib по умолчанию не доверяет IP из ответа PASV и открывает канал
    # данных на тот же адрес, что и управляющее соединение — оставляем так.
    target_ip = validate_outbound_target(host, port)
    ftp = ftplib.FTP_TLS() if use_tls else ftplib.FTP()
    ftp.connect(host=target_ip, port=port or 21, timeout=20)
    ftp.login(user=user or "", passwd=password or "")
    if use_tls:
        ftp.prot_p()
    return ftp


def download_ftp_file(host, port, user, password, remote_path, use_tls=False) -> bytes:
    ftp = _connect_ftp(host, port, user, password, use_tls)
    try:
        buffer = io.BytesIO()
        started = time.monotonic()

        def _write(chunk: bytes) -> None:
            if buffer.tell() + len(chunk) > MAX_FTP_FILE_BYTES:
                raise ValueError(
                    f"Файл на FTP больше допустимого размера ({MAX_FTP_FILE_BYTES // (1024 * 1024)} МБ)"
                )
            if time.monotonic() - started > MAX_FTP_DOWNLOAD_SECONDS:
                raise TimeoutError("Скачивание файла с FTP заняло слишком много времени")
            buffer.write(chunk)

        ftp.retrbinary(f"RETR {remote_path}", _write)
        return buffer.getvalue()
    finally:
        try:
            ftp.quit()
        except Exception:
            ftp.close()


def fetch_price_file(host, port, user, password, remote_path, use_tls=False):
    """(имя файла, байты). Адрес «local» — встроенная папка прайсов, которую
    наполняет собственный FTP-сервер (см. price_inbox.py); иначе — внешний FTP."""
    if price_inbox.is_local_host(host):
        return price_inbox.read_file(remote_path, MAX_FTP_FILE_BYTES)
    return remote_path, download_ftp_file(host, port, user, password, remote_path, use_tls)


def test_ftp_connection(host, port, user, password, remote_path, use_tls=False) -> dict:
    """Селфчек: подключается и проверяет, что указанный файл виден и
    доступен, не скачивая и не парся его целиком. Используется и отдельной
    кнопкой "Проверить соединение" в админке, и общим self-check'ом."""
    if price_inbox.is_local_host(host):
        try:
            info = price_inbox.describe_file(remote_path)
        except price_inbox.InboxError as exc:
            return {"ok": False, "message": str(exc)}
        return {
            "ok": True,
            "message": f"Встроенная папка прайсов: найден файл «{info['name']}»",
            "size_bytes": info["size_bytes"],
            "modified_at": info["modified_at"],
        }
    if not host or not remote_path:
        return {"ok": False, "message": "Не заданы адрес сервера или путь к файлу"}
    try:
        ftp = _connect_ftp(host, port, user, password, use_tls)
    except Exception as exc:
        return {"ok": False, "message": f"Не удалось подключиться: {exc}"}

    try:
        size = None
        mtime = None
        try:
            size = ftp.size(remote_path)
        except Exception:
            pass
        try:
            mtime = ftp.voidcmd(f"MDTM {remote_path}")[4:].strip()
        except Exception:
            pass

        if size is None and mtime is None:
            # Сервер не поддерживает SIZE/MDTM — проверяем через список каталога.
            directory, _, filename = remote_path.rpartition("/")
            try:
                names = ftp.nlst(directory or ".")
            except Exception as exc:
                return {"ok": False, "message": f"Подключение успешно, но файл не проверить: {exc}"}
            if not any(filename and filename in n for n in names):
                return {"ok": False, "message": f"Подключение успешно, но файл {remote_path} не найден"}

        return {"ok": True, "message": "Подключение и файл в порядке", "size_bytes": size, "modified_at": mtime}
    finally:
        try:
            ftp.quit()
        except Exception:
            ftp.close()


def _build_oem_code_index(db) -> dict:
    """code (lower-case, без пробелов по краям) -> список part_id, у которых
    в атрибуте #9048 (ОЕМ-номер) есть ровно такое значение среди "; "-
    разделённых. Один код может относиться к нескольким part_id (см. верхний
    docstring про неоднозначные совпадения) — это нормально, обрабатывается
    в run_ftp_price_sync."""
    index = defaultdict(list)
    rows = (
        db.query(PartAttribute.part_id, PartAttribute.value)
        .filter(PartAttribute.ozon_attribute_id == ARTICLE_ATTRIBUTE_ID)
        .filter(PartAttribute.value.isnot(None))
        .all()
    )
    for part_id, value in rows:
        for code in value.split(";"):
            code = code.strip().lower()
            if code:
                index[code].append(part_id)
    return index


async def run_ftp_price_sync(trigger: str = "scheduled"):
    init_db()
    db = SessionLocal()
    settings = get_or_create_sync_settings(db)

    log_entry = SyncLog(source="ftp", status="running")
    db.add(log_entry)
    db.commit()

    matched = 0
    unmatched = 0
    unmatched_sample = []

    try:
        local = price_inbox.is_local_host(settings.ftp_host)
        if not settings.ftp_host or (not local and not settings.ftp_remote_path):
            raise RuntimeError("FTP не настроен: не задан адрес сервера или путь к файлу")

        file_name, raw_bytes = fetch_price_file(
            settings.ftp_host, settings.ftp_port, settings.ftp_user,
            settings.ftp_password, settings.ftp_remote_path, settings.ftp_use_tls,
        )
        rows = parse_price_table(file_name, raw_bytes)
        if not rows:
            raise RuntimeError("Файл скачан, но не удалось распознать ни одной строки «артикул — цена»")

        code_to_part_ids = _build_oem_code_index(db)

        ambiguous = 0
        ambiguous_sample = []
        seen_part_ids = set()
        for article, price in rows:
            key = article.strip().lower()
            part_ids = code_to_part_ids.get(key) or []
            if not part_ids:
                unmatched += 1
                if len(unmatched_sample) < MAX_UNMATCHED_SAMPLE:
                    unmatched_sample.append(article)
                continue
            if len(part_ids) > 1:
                # Один и тот же ОЕМ-номер указан в карточках нескольких
                # разных товаров — какому из них принадлежит цена, неясно,
                # поэтому никому её не проставляем.
                ambiguous += 1
                if len(ambiguous_sample) < MAX_UNMATCHED_SAMPLE:
                    ambiguous_sample.append(article)
                continue
            part = db.get(Part, part_ids[0])
            if part is None:
                unmatched += 1
                continue
            part.ftp_price = price
            part.ftp_price_updated_at = _now()
            part.price = price
            seen_part_ids.add(part.id)
            matched += 1

        # Товары, у которых раньше была цена с FTP, но в этом файле их больше
        # нет — снимаем приоритет, дальше цену им снова ведёт Ozon.
        stale_query = db.query(Part).filter(Part.ftp_price.isnot(None))
        if seen_part_ids:
            stale_query = stale_query.filter(~Part.id.in_(seen_part_ids))
        for part in stale_query.all():
            part.ftp_price = None
            part.ftp_price_updated_at = None

        db.commit()
        log_entry.status = "success" if unmatched == 0 and ambiguous == 0 else "success_with_errors"
    except Exception as exc:
        db.rollback()
        log_entry.status = "failed"
        log_entry.error_message = str(exc)
        logger.exception("Импорт цен с FTP завершился с ошибкой: %s", exc)
    finally:
        log_entry.finished_at = _now()
        log_entry.products_processed = matched
        log_entry.products_failed = unmatched
        log_entry.details = {
            "trigger": trigger,
            "matched": matched,
            "unmatched": unmatched,
            "unmatched_sample": unmatched_sample,
            "ambiguous": ambiguous if "ambiguous" in locals() else 0,
            "ambiguous_sample": ambiguous_sample if "ambiguous_sample" in locals() else [],
        }
        db.commit()
        db.close()
        logger.info(
            "Импорт цен с FTP завершён: совпало=%d, не найдено=%d, неоднозначно=%d",
            matched, unmatched, ambiguous if "ambiguous" in locals() else 0,
        )
